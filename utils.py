import re
import io
import streamlit as st
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from pdfminer.high_level import extract_text
from docx import Document

def extract_text_from_pdf(uploaded_file):
    """Extracts text safely from an uploaded PDF file object in-memory."""
    try:
        if hasattr(uploaded_file, 'getvalue'):
            pdf_stream = io.BytesIO(uploaded_file.getvalue())
            return extract_text(pdf_stream)
        return extract_text(uploaded_file)
    except Exception as e:
        st.error(f"Error parsing PDF file: {e}")
        return ""

def extract_text_from_docx(uploaded_file):
    """Extracts text safely from an uploaded DOCX file object in-memory."""
    try:
        if hasattr(uploaded_file, 'getvalue'):
            docx_stream = io.BytesIO(uploaded_file.getvalue())
            doc = Document(docx_stream)
        else:
            doc = Document(uploaded_file)
            
        full_text = []
        # Extract structural text from standard paragraphs
        for paragraph in doc.paragraphs:
            full_text.append(paragraph.text)
            
        # Extract text out of tables (vital for resumes styled inside table matrices)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
                    
        return "\n".join(full_text)
    except Exception as e:
        st.error(f"Error parsing DOCX file: {e}")
        return ""

def clean_text(text):
    """Normalizes text by lowercase conversion, space reduction, and punctuation removal."""
    if not text:
        return ""
    text = text.lower()
    text = re.sub(r'\s+', ' ', text)  # Normalizes newlines and multiple spaces
    text = re.sub(r'[^a-zA-Z0-9\s]', '', text)
    return text.strip()

def calculate_similarity(resume_text, jd_text):
    """Calculates TF-IDF cosine similarity percentage between text segments, ignoring stop words."""
    if not resume_text or not jd_text:
        return 0.0
    try:
        vectorizer = TfidfVectorizer(stop_words='english')
        vectors = vectorizer.fit_transform([resume_text, jd_text])
        score = cosine_similarity(vectors[0:1], vectors[1:2])
        return round(float(score) * 100, 2)
    except ValueError:
        return 0.0

def find_missing_keywords(resume_clean, jd_clean):
    """Identifies distinct words >5 characters present in the job description but omitted from the resume."""
    jd_words = set(word for word in jd_clean.split() if len(word) > 5)
    resume_words = set(resume_clean.split())
    missing = list(jd_words - resume_words)
    return ", ".join(sorted(missing)[:5])
